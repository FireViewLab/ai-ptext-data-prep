"""Build an auditable binary dataset; model input is content only, never provenance."""

import argparse
from collections import Counter, defaultdict
import hashlib
import json
import os
from pathlib import Path
import random
import sys
import tempfile

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from generate_ptext_synthetic import TYPES, generate
from inspect_reviews import PROJECT_ROOT, is_blank, markdown_value, read_reviews


COLUMNS = (
    "platform", "product_id", "review_id", "content", "label", "label_origin",
    "synthetic_type", "parent_review_id", "group_id", "split", "source_excel_row",
    "parent_source_excel_row", "synthetic_family_id", "generation_method", "seed", "evaluation_set",
)
SPLITS = ("train", "validation", "test")
RATIOS = np.array([.70, .15, .15])


def fingerprint(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalized(content):
    return " ".join(content.split())


class Groups:
    def __init__(self, size):
        self.parent = list(range(size))

    def find(self, index):
        while self.parent[index] != index:
            self.parent[index] = self.parent[self.parent[index]]
            index = self.parent[index]
        return index

    def union(self, left, right):
        a, b = sorted((self.find(left), self.find(right)))
        self.parent[b] = a


def union_equal(groups, entries):
    seen = {}
    for index, value in entries:
        if value in seen:
            groups.union(index, seen[value])
        else:
            seen[value] = index


def union_similar(groups, contents, threshold=.85):
    """Group exact text at every length; char-wb cosine for length >= 15, in batches."""
    union_equal(groups, ((i, normalized(text)) for i, text in enumerate(contents) if text))
    positions = [i for i, text in enumerate(contents) if len(normalized(text)) >= 15]
    if len(positions) < 2:
        return 0
    texts = [normalized(contents[i]) for i in positions]
    matrix = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5), lowercase=False).fit_transform(texts)
    edges = 0
    for start in range(0, len(texts), 64):
        scores = cosine_similarity(matrix[start:start + 64], matrix, dense_output=False).tocsr()
        for local in range(scores.shape[0]):
            left = start + local
            begin, end = scores.indptr[local:local + 2]
            qualifying = np.flatnonzero(scores.data[begin:end] >= threshold - 1e-12)
            for relative_offset in qualifying:
                offset = begin + relative_offset
                right = int(scores.indices[offset])
                if right > left:
                    groups.union(positions[left], positions[right])
                    edges += 1
    return edges


def read_inputs(review_path, label_path):
    for path in (review_path, label_path):
        if not path.is_file() or path.stat().st_size == 0:
            raise ValueError(f"입력 파일이 없거나 0바이트입니다: {path}")
    frame, _, _, _ = read_reviews(review_path)
    raw = [dict(zip(["source_excel_row", *frame.columns], values))
           for values in frame.itertuples(index=True, name=None)]
    by_row = {row["source_excel_row"]: row for row in raw}
    workbook = load_workbook(label_path, read_only=True, data_only=False)
    labels = {}
    try:
        if "candidates" not in workbook.sheetnames:
            raise ValueError("사람 검토 파일에 candidates 시트가 없습니다.")
        iterator = workbook["candidates"].iter_rows(values_only=True)
        header = list(next(iterator, ()))
        required = ("source_excel_row", "platform", "product_id", "review_id", "content", "human_label")
        if any(header.count(name) != 1 for name in required):
            raise ValueError("사람 검토 파일의 필수 헤더가 없거나 중복되었습니다.")
        for values in iterator:
            if all(is_blank(value) for value in values):
                continue
            row = dict(zip(header, values))
            number = row["source_excel_row"]
            if number not in by_row or number in labels:
                raise ValueError(f"사람 검토 원본 행 번호가 없거나 중복되었습니다: {number}")
            for field in ("platform", "product_id", "review_id", "content"):
                left, right = row[field], by_row[number][field]
                if not ((is_blank(left) and is_blank(right)) or left == right):
                    raise ValueError(f"사람 검토 원본 대조 실패: Excel {number}행 {field}")
            if row["human_label"] not in ("NORMAL", "SUSPICIOUS", "UNCERTAIN"):
                raise ValueError(f"유효하지 않거나 비어 있는 human_label: Excel {number}행")
            labels[number] = row["human_label"]
    finally:
        workbook.close()
    if not labels:
        raise ValueError("검토 완료 라벨이 없습니다.")
    return raw, labels


def stratified_sample(pool, count, seed, musinsa_cap=20):
    """Proportional highest-averages allocation; explicitly cap scarce musinsa."""
    rng = random.Random(seed)
    strata = defaultdict(list)
    for row in pool:
        strata[row["platform"]].append(row)
    capacities = {key: min(len(rows), musinsa_cap) if key == "musinsa" else len(rows)
                  for key, rows in strata.items()}
    target = min(count, sum(capacities.values()))
    quotas = {key: 0 for key in strata}
    for _ in range(target):
        eligible = [key for key in strata if quotas[key] < capacities[key]]
        key = max(sorted(eligible, key=str), key=lambda key: len(strata[key]) / (quotas[key] + 1))
        quotas[key] += 1
    chosen = []
    for key in sorted(strata, key=str):
        chosen.extend(rng.sample(strata[key], quotas[key]))
    return sorted(chosen, key=lambda row: row["source_excel_row"]), quotas


def assign_splits(rows, seed):
    """Greedy group-level balancing of counts, origin, label and platform."""
    groups = defaultdict(list)
    for row in rows:
        groups[row["group_id"]].append(row)
    def features(row):
        return ("all", "origin:" + row["label_origin"], "label:" + row["label"], "platform:" + str(row["platform"]))
    dimensions = sorted({feature for row in rows for feature in features(row)})
    index = {value: i for i, value in enumerate(dimensions)}
    vectors = {}
    for group, members in groups.items():
        vector = np.zeros(len(dimensions))
        for row in members:
            for feature in features(row):
                vector[index[feature]] += 1
        vectors[group] = vector
    totals = sum(vectors.values(), np.zeros(len(dimensions)))
    targets = RATIOS[:, None] * totals[None, :]
    current = np.zeros_like(targets)
    rng = random.Random(seed)
    order = list(groups)
    rng.shuffle(order)
    order.sort(key=lambda group: -len(groups[group]))
    for group in order:
        vector = vectors[group]
        before = (current - targets) ** 2 / np.maximum(targets, 1)
        after = (current + vector - targets) ** 2 / np.maximum(targets, 1)
        split_index = int(np.argmin((after - before).sum(axis=1)))
        current[split_index] += vector
        for row in groups[group]:
            row["split"] = SPLITS[split_index]
            row["evaluation_set"] = (
                "training" if split_index == 0 else
                "synthetic_challenge" if row["label_origin"] == "synthetic" else
                "real_human_evaluation" if row["label_origin"] in ("human_candidate_review", "human_suspicious") else
                "weak_label_evaluation"
            )


def validate(rows, labels):
    for row in rows:
        if not isinstance(row["content"], str) or is_blank(row["content"]):
            raise ValueError("빈 content 또는 비문자열 content가 학습 행에 있습니다.")
        if row["label"] not in ("NORMAL", "SUSPICIOUS") or row["split"] not in SPLITS:
            raise ValueError("허용되지 않은 label/split입니다.")
        synthetic = row["label_origin"] == "synthetic"
        if synthetic and row["synthetic_type"] not in TYPES:
            raise ValueError("synthetic_type 누락 또는 오류")
        if not synthetic and row.get("synthetic_type"):
            raise ValueError("실제 리뷰에 synthetic_type이 있습니다.")
        if labels.get(row.get("source_excel_row")) == "UNCERTAIN":
            raise ValueError("UNCERTAIN 행이 학습 데이터에 포함되었습니다.")
    leakage = {}
    for field in ("group_id", "parent_review_id", "synthetic_family_id"):
        buckets = defaultdict(set)
        for row in rows:
            value = row.get(field)
            if not is_blank(value):
                buckets[value].add(row["split"])
        leakage[field] = sum(len(splits) > 1 for splits in buckets.values())
    real = {row["source_excel_row"]: row for row in rows if row["label_origin"] != "synthetic"}
    parent_failures = 0
    for row in rows:
        if row["label_origin"] == "synthetic":
            parent = real.get(row["parent_source_excel_row"])
            if parent is None or parent["split"] != row["split"] or parent["group_id"] != row["group_id"] or parent["review_id"] != row["parent_review_id"]:
                parent_failures += 1
    leakage["parent_source_row"] = parent_failures
    if any(leakage.values()):
        raise ValueError(f"leakage 검증 실패: {leakage}")
    return leakage


def summarize(rows):
    result = {"total": len(rows)}
    for field in ("label", "label_origin", "synthetic_type", "platform", "split", "evaluation_set"):
        result[field] = dict(Counter(row.get(field) or "[real/empty]" for row in rows))
    for fields in (("platform", "label"), ("split", "label"), ("split", "label_origin"), ("evaluation_set", "split", "label")):
        result["_x_".join(fields)] = dict(Counter(" | ".join(str(row[field]) for field in fields) for row in rows))
    return result


def build_dataset(review_path, label_path, sample_count=1500, synthetic_count=1350, seed=42, near_threshold=.85):
    if min(sample_count, synthetic_count) < 0 or not 0 < near_threshold <= 1:
        raise ValueError("샘플 수는 0 이상, near threshold는 0 초과 1 이하여야 합니다.")
    raw, labels = read_inputs(review_path, label_path)
    valid = [row for row in raw if isinstance(row["content"], str) and not is_blank(row["content"])]
    valid_numbers = {row["source_excel_row"] for row in valid}
    if any(number not in valid_numbers and label != "UNCERTAIN" for number, label in labels.items()):
        raise ValueError("사람 검토 확정 행의 content가 비어 있습니다. 자동 제외하지 않습니다.")
    groups = Groups(len(valid))
    real_edges = union_similar(groups, [row["content"] for row in valid], near_threshold)
    # Conservatively group repeated bare IDs too, so parent_review_id is unambiguous across splits.
    union_equal(groups, ((i, row["review_id"]) for i, row in enumerate(valid) if not is_blank(row["review_id"])))
    reviewed_groups = {groups.find(i) for i, row in enumerate(valid) if row["source_excel_row"] in labels}
    pool = [row for i, row in enumerate(valid) if row["source_excel_row"] not in labels and groups.find(i) not in reviewed_groups]
    sampled, quotas = stratified_sample(pool, sample_count, seed)
    rows = []
    for source in valid:
        label = labels.get(source["source_excel_row"])
        if label and label != "UNCERTAIN":
            rows.append({**source, "label": label, "label_origin": "human_suspicious" if label == "SUSPICIOUS" else "human_candidate_review"})
    rows += [{**row, "label": "NORMAL", "label_origin": "sampled_non_candidate"} for row in sampled]
    for row in rows:
        row.update(synthetic_type=None, parent_review_id=None, parent_source_excel_row=None,
                   synthetic_family_id=None, generation_method=None, seed=seed)
    uncertain_groups = {groups.find(i) for i, row in enumerate(valid) if labels.get(row["source_excel_row"]) == "UNCERTAIN"}
    original_index = {row["source_excel_row"]: i for i, row in enumerate(valid)}
    parents = [row for row in rows if not is_blank(row["review_id"])
               and len(normalized(row["content"])) >= 15
               and groups.find(original_index[row["source_excel_row"]]) not in uncertain_groups]
    synthetic = generate(parents, synthetic_count, seed)
    rows.extend(synthetic)
    # Compute final text similarity including synthetic variants. All original
    # reviews remain in this graph so excluded intermediates cannot split a family.
    final_groups = Groups(len(valid) + len(synthetic))
    for i in range(len(valid)):
        final_groups.union(i, groups.find(i))
    for i, row in enumerate(synthetic, len(valid)):
        final_groups.union(i, original_index[row["parent_source_excel_row"]])
    final_edges = union_similar(final_groups, [row["content"] for row in valid + synthetic], near_threshold)
    for row in rows:
        index = original_index[row["source_excel_row"]] if row["label_origin"] != "synthetic" else original_index[row["parent_source_excel_row"]]
        row["group_id"] = f"group-{valid[final_groups.find(index)]['source_excel_row']:06d}"
    assign_splits(rows, seed)
    leakage = validate(rows, labels)
    metadata = {
        "seed": seed, "input_reviews": len(raw), "human_labels": dict(Counter(labels.values())),
        "input_platform_counts": dict(Counter(row["platform"] for row in raw)),
        "uncertain_excluded": sum(label == "UNCERTAIN" for label in labels.values()),
        "sample_requested": sample_count, "sample_actual": len(sampled), "sample_platforms": quotas,
        "eligible_non_candidate_pool": len(pool), "near_threshold": near_threshold,
        "real_near_edges": real_edges, "final_near_edges": final_edges, "leakage": leakage,
        "largest_group": max(Counter(row["group_id"] for row in rows).values(), default=0),
        "human_normal_split": dict(Counter(row["split"] for row in rows if row["label_origin"] == "human_candidate_review")),
        "synthetic_unique_texts": len({row["content"] for row in synthetic}),
        "synthetic_length_min_max": [min(map(lambda r: len(r['content']), synthetic), default=0), max(map(lambda r: len(r['content']), synthetic), default=0)],
    }
    return rows, summarize(rows), metadata


def report_text(summary, metadata, hashes):
    lines = ["# P_text training dataset v1", "",
             "모델 입력은 content 한 컬럼만 사용합니다. label_origin, synthetic_type, ID, platform 등은 모델 입력에서 제외합니다.",
             "실제 사람 검토 라벨 / 비후보 추정 NORMAL / synthetic SUSPICIOUS를 별도로 기록합니다.", ""]
    for field, values in summary.items():
        lines += [f"## {field}", "", "| 항목 | 수 |", "| --- | ---: |"]
        for name, count in (values.items() if isinstance(values, dict) else [(field, values)]):
            lines.append(f"| {markdown_value(name)} | {count} |")
        lines.append("")
    lines += ["## 검증 및 생성 정보", "", "```json", json.dumps(metadata, ensure_ascii=False, indent=2), "```", "",
              "## 평가와 한계", "",
              "- 그룹 단위 70/15/15 목표 분할입니다. 연결된 원문/부모/유사문장 그룹은 쪼개지 않으므로 실제 비율은 달라질 수 있습니다.",
              "- 모든 실제 원문의 공백 정리 후 exact 중복 및 길이 15자 이상 char_wb 3~5-gram cosine 유사도 기준 연결요소를 그룹화합니다. 최종 synthetic까지 포함해 다시 검사합니다. 이 기준 이하의 의미상 유사성까지 보장하지 않습니다.",
              "- 사람이 검토한 NORMAL은 전부 유지합니다. 비후보 샘플은 검토 후보와 같은 유사 그룹도 제외하고 platform 비례 샘플링합니다. musinsa는 최대 20건만 추가 샘플링하며 데이터 자체가 적습니다.",
              "- sampled_non_candidate는 사람이 확인한 NORMAL이 아닌 약한 추정 라벨입니다. 별도 수동 확인이 필요합니다.",
              "- Real-data evaluation set: validation/test의 human_candidate_review 및 human_suspicious만 평가합니다. weak_label_evaluation은 별도로 보고하고 확정 라벨 성능과 섞지 않습니다.",
              "- Synthetic challenge set: validation/test의 synthetic 행을 별도로 평가합니다. 실제 조작 리뷰 성능의 대체 지표로 사용하지 않습니다.",
              "- 실제 SUSPICIOUS가 현재 1건뿐이므로 실제 positive recall을 신뢰성 있게 측정할 수 없습니다. 그 행이 train에 배정되면 실제 평가셋의 positive recall은 정의할 수 없습니다.",
              "- synthetic은 원문에서 확인되는 상품/속성 어휘를 이용한 재현 가능한 조합형 생성입니다. 사람이나 외부 생성 모델의 검수를 거친 자연어가 아니며 부자연스러운 문장과 공통 문체가 남을 수 있습니다.",
              "- synthetic의 효능·성능 단정은 인위적으로 만든 주장입니다. 실제 제품의 확인된 사실이 아닙니다. synthetic_type별 균형, 고유 본문 및 길이 범위를 기록하지만 의미적 다양성이나 품질을 보장하지 않습니다.",
              "- COPY_VARIATION은 같은 홍보성 prototype의 어휘/문장 순서 변형이며 동일 family와 원본 부모 그룹에 묶습니다. 리뷰 간 유사성 자체는 P_text 위험 라벨 근거가 아닙니다. 각 변형의 경험 없는 홍보 본문을 인위적 challenge로 사용하는 것입니다.",
              "- 자동 생성 레이블을 실제 사람이 확인한 레이블로 간주하지 마세요. 모델 학습은 수행하지 않았습니다.", "",
              "## 입력 원본 SHA256", ""]
    for path, value in hashes.items():
        lines.append(f"- {markdown_value(path)}: `{value}`")
    test_total = summary["split"].get("test", 0)
    synth_test = summary["split_x_label_origin"].get("test | synthetic", 0)
    lines += ["", f"Test synthetic 비율: {synth_test}/{test_total} ({synth_test / test_total:.2%})" if test_total else "Test 행 없음: 평가 불가", ""]
    return "\n".join(lines)


def save_outputs(rows, summary, metadata, hashes, output, report):
    temporary = []
    try:
        for path in (output, report):
            if path.exists():
                raise ValueError(f"결과 파일을 덮어쓰지 않습니다: {path}")
            path.parent.mkdir(parents=True, exist_ok=True)
        for path in (output, report):
            with tempfile.NamedTemporaryFile(mode="w+b", dir=path.parent, prefix=path.stem + ".", suffix=".tmp", delete=False) as stream:
                temporary.append(Path(stream.name))
                if path == output:
                    workbook = Workbook(write_only=True)
                    try:
                        sheet = workbook.create_sheet("dataset")
                        sheet.append(list(COLUMNS))
                        for row in rows:
                            cells = []
                            for name in COLUMNS:
                                value = row.get(name)
                                cell = WriteOnlyCell(sheet, value=value)
                                if isinstance(value, str):
                                    cell.data_type = "s"
                                cells.append(cell)
                            sheet.append(cells)
                        sheet = workbook.create_sheet("summary")
                        sheet.append(["section", "item", "value"])
                        for name, values in summary.items():
                            for item, value in (values.items() if isinstance(values, dict) else [(name, values)]):
                                sheet.append([name, str(item), value])
                        workbook.save(stream)
                    finally:
                        workbook.close()
                else:
                    stream.write(report_text(summary, metadata, hashes).encode("utf-8"))
                stream.flush()
                os.fsync(stream.fileno())
        for tmp, path in zip(temporary, (output, report)):
            if os.name == "nt":
                tmp.rename(path)
            else:
                os.link(tmp, path)
    finally:
        for path in temporary:
            path.unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser(description="P_text 최종 데이터셋 구성 (모델 학습 없음)")
    parser.add_argument("--reviews", type=Path, default=Path("outputs/merged_reviews.xlsx"))
    parser.add_argument("--labels", type=Path, default=Path("outputs/merged_ptext_candidates_human_review_labeled_v1.xlsx"))
    parser.add_argument("--output", type=Path, default=Path("outputs/ptext_training_dataset_v1.xlsx"))
    parser.add_argument("--report", type=Path, default=Path("reports/ptext_training_dataset_v1_summary.md"))
    parser.add_argument("--sample-count", type=int, default=1500)
    parser.add_argument("--synthetic-count", type=int, default=1350)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--near-threshold", type=float, default=.85)
    args = parser.parse_args()
    paths = [(PROJECT_ROOT / path.expanduser()).resolve() for path in (args.reviews, args.labels, args.output, args.report)]
    reviews, labels, output, report = paths
    try:
        if output == report or any(path in (reviews, labels) for path in (output, report)):
            raise ValueError("입력/출력 경로가 충돌합니다.")
        if output.exists() or report.exists():
            raise ValueError("기존 결과를 덮어쓰지 않습니다.")
        before = {path: fingerprint(path) for path in (reviews, labels)}
        rows, summary, metadata = build_dataset(reviews, labels, args.sample_count, args.synthetic_count, args.seed, args.near_threshold)
        if any(fingerprint(path) != value for path, value in before.items()):
            raise ValueError("입력 원본이 실행 중 변경되었습니다. 결과를 저장하지 않습니다.")
        hashes = {str(path.relative_to(PROJECT_ROOT)) if path.is_relative_to(PROJECT_ROOT) else str(path): value for path, value in before.items()}
        save_outputs(rows, summary, metadata, hashes, output, report)
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        print(f"저장 완료: {output}\n보고서: {report}")
    except Exception as exc:
        print(f"오류: 데이터셋 생성 실패: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
