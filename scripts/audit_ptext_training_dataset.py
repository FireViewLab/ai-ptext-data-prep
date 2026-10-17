"""Read-only dataset audit with deterministic review samples; writes Markdown only."""

import argparse
from collections import Counter, defaultdict
import hashlib
import html
import json
from pathlib import Path
import random
import re
from statistics import mean, median

from openpyxl import load_workbook

from build_ptext_training_dataset import PROJECT_ROOT, SPLITS, validate
from generate_ptext_synthetic import TYPES


def read_dataset(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        iterator = workbook["dataset"].iter_rows(values_only=True)
        header = next(iterator)
        required = ("content", "label", "label_origin", "synthetic_type", "split",
                    "group_id", "parent_review_id", "parent_source_excel_row",
                    "synthetic_family_id", "source_excel_row", "review_id", "platform")
        if any(header.count(column) != 1 for column in required):
            raise ValueError("dataset 필수 컬럼 누락 또는 중복")
        return [{**dict(zip(header, values)), "dataset_excel_row": number}
                for number, values in enumerate(iterator, start=2)]
    finally:
        workbook.close()


def exact_stats(rows):
    """Exact means original string equality, including all spaces and newlines."""
    counts = Counter(row["content"] for row in rows)
    repeated = [count for count in counts.values() if count > 1]
    return {"groups": len(repeated), "participating_rows": sum(repeated),
            "extra_rows": sum(count - 1 for count in repeated)}


def conflicts(rows):
    groups = defaultdict(list)
    for row in rows:
        groups[row["content"]].append(row)
    return [members for members in groups.values() if len({row["label"] for row in members}) > 1]


def lengths(rows):
    values = [len(row["content"]) for row in rows]
    return {"count": len(values), "mean": round(mean(values), 2) if values else 0,
            "median": median(values) if values else 0, "min": min(values, default=0),
            "max": max(values, default=0)}


def document_phrases(rows):
    """Document frequency of 2-5 whitespace-token phrases, diagnostic copies only."""
    counts = Counter()
    for row in rows:
        words = re.findall(r"[가-힣A-Za-z0-9]+", row["content"])
        phrases = {" ".join(words[i:i + size]) for size in range(2, 6)
                   for i in range(len(words) - size + 1)}
        counts.update(phrase for phrase in phrases if len(phrase) >= 6)
    # Prefer longer equal-frequency phrases over their contained fragments.
    selected = []
    for phrase, count in sorted(counts.items(), key=lambda item: (-item[1], -len(item[0]), item[0])):
        if count < 2:
            break
        if not any(phrase in longer and count == earlier_count for longer, earlier_count in selected):
            selected.append((phrase, count))
        if len(selected) == 20:
            break
    return selected


def deterministic_samples(rows, seed=42):
    rng = random.Random(seed)
    categories = [(kind, [r for r in rows if r["label_origin"] == "synthetic" and r["synthetic_type"] == kind], 10)
                  for kind in TYPES]
    categories += [
        ("sampled_non_candidate NORMAL", [r for r in rows if r["label_origin"] == "sampled_non_candidate" and r["label"] == "NORMAL"], 30),
        ("human_candidate_review NORMAL", [r for r in rows if r["label_origin"] == "human_candidate_review" and r["label"] == "NORMAL"], 20),
        ("human SUSPICIOUS", [r for r in rows if r["label_origin"] != "synthetic" and r["label_origin"] != "sampled_non_candidate" and r["label"] == "SUSPICIOUS"], 1),
    ]
    return {name: rng.sample(sorted(pool, key=lambda row: row["dataset_excel_row"]), min(count, len(pool)))
            for name, pool, count in categories}


def leakage_stats(rows):
    result = {}
    for field in ("group_id", "parent_review_id", "synthetic_family_id"):
        buckets = defaultdict(set)
        for row in rows:
            if row.get(field) is not None and row.get(field) != "":
                buckets[row[field]].add(row["split"])
        result[field] = sum(len(splits) > 1 for splits in buckets.values())
    by_source = {r["source_excel_row"]: r for r in rows if r["label_origin"] != "synthetic"}
    result["parent_source_row"] = 0
    for row in rows:
        if row["label_origin"] == "synthetic":
            parent = by_source.get(row["parent_source_excel_row"])
            if parent is None or any(parent[key] != row[key] for key in ("split", "group_id")) or parent["review_id"] != row["parent_review_id"]:
                result["parent_source_row"] += 1
    try:
        validate(rows, {})
        result["existing_validator"] = "PASS"
    except ValueError as exc:
        result["existing_validator"] = str(exc)
    return result


def audit(rows, seed=42):
    synthetic = [r for r in rows if r["label_origin"] == "synthetic"]
    suspicious = [r for r in rows if r["label"] == "SUSPICIOUS"]
    normal = [r for r in rows if r["label"] == "NORMAL"]
    repeated = conflicts(rows)
    # Prespecified lexical diagnostics: document presence, not occurrence count.
    probes = ("선택", "구성", "확실", "보장", "기대", "품질", "만족", "추천", "구매", "제품", "관리", "중심", "분들을 위한")
    result = {
        "total": len(rows), "conflict_groups": len(repeated),
        "conflict_rows": sum(len(group) for group in repeated),
        "conflict_pairs": sum(sum(r["label"] == "NORMAL" for r in group) * sum(r["label"] == "SUSPICIOUS" for r in group) for group in repeated),
        "exact": {"all": exact_stats(rows)},
        "lengths": {}, "starts": {}, "expressions": {},
        "label_origin_x_split": dict(Counter((r["label_origin"], r["split"]) for r in rows)),
        "platform_x_label_x_split": dict(Counter((r["platform"], r["label"], r["split"]) for r in rows)),
        "leakage": leakage_stats(rows),
        "probes": [(phrase, sum(phrase in r["content"] for r in suspicious), len(suspicious),
                    sum(phrase in r["content"] for r in normal), len(normal)) for phrase in probes],
    }
    for label in ("NORMAL", "SUSPICIOUS"):
        result["exact"][label] = exact_stats([r for r in rows if r["label"] == label])
    for kind in TYPES:
        members = [r for r in synthetic if r["synthetic_type"] == kind]
        result["exact"][kind] = exact_stats(members)
        result["lengths"][kind] = lengths(members)
        result["starts"][kind] = Counter(" ".join(r["content"].split()[:4]) for r in members).most_common(10)
        result["expressions"][kind] = document_phrases(members)
    result["starts"]["all synthetic"] = Counter(" ".join(r["content"].split()[:4]) for r in synthetic).most_common(20)
    result["expressions"]["all synthetic"] = document_phrases(synthetic)
    result["lengths"]["real NORMAL"] = lengths([r for r in normal if r["label_origin"] != "synthetic"])
    result["lengths"]["all synthetic"] = lengths(synthetic)
    return result, repeated, deterministic_samples(rows, seed)


def escaped(value):
    return html.escape(str(value)).replace("|", "&#124;").replace("\n", "\\n").replace("\r", "\\r")


def table(headers, data):
    return ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |",
            *("| " + " | ".join(escaped(value) for value in row) + " |" for row in data), ""]


def render_report(rows, result, repeated, samples, input_hash, seed):
    lines = ["# P_text training dataset v1 품질 audit", "",
             f"- 입력: outputs/ptext_training_dataset_v1.xlsx / dataset ({len(rows)}행)",
             f"- 입력 SHA256: `{input_hash}`", f"- 무작위 샘플 seed: {seed}",
             "- 데이터셋·라벨은 읽기 전용으로 검사했습니다. 정제, 재라벨링, 학습을 수행하지 않았습니다.", "",
             "## 1. 동일 원문 label 충돌", "",
             f"충돌 content 그룹 {result['conflict_groups']}개 / 참여 행 {result['conflict_rows']}개 / NORMAL×SUSPICIOUS pair {result['conflict_pairs']}개.", ""]
    for index, group in enumerate(repeated, start=1):
        lines += [f"충돌 그룹 {index} (본문 미기재; 행 번호는 dataset 시트의 Excel 행)", ""]
        lines += table(["Excel 행", "label_origin"], [(r["dataset_excel_row"], r["label_origin"]) for r in group])
    lines += ["## 2. Exact duplicate", "",
              "원문 공백·줄바꿈을 포함한 정확한 문자열 동등성 기준입니다. 그룹 수 / 최초 행 포함 참여 행 / 최초 1행을 뺀 추가 행을 구분합니다. 타입별 집계는 해당 타입 내부 중복입니다.", ""]
    lines += table(["범위", "중복 그룹", "참여 행", "추가 중복 행"],
                   [(name, stats["groups"], stats["participating_rows"], stats["extra_rows"]) for name, stats in result["exact"].items()])
    lines += ["## 3. Synthetic 다양성", "", "길이는 원본 문자열의 문자 수입니다.", ""]
    lines += table(["유형", "행 수", "평균", "중앙값", "최소", "최대"],
                   [(name, *(stats[key] for key in ("count", "mean", "median", "min", "max"))) for name, stats in result["lengths"].items()])
    lines += ["### 시작 구문", "",
              "진단용으로만 공백을 정리한 첫 4개 어절을 집계합니다. 시작 4어절이 같다고 전체 원문이 동일하다는 뜻은 아닙니다. 유형 내 비율 10% 이상은 집중 검토 대상으로 표시합니다.", ""]
    for kind, entries in result["starts"].items():
        denominator = result["lengths"][kind]["count"]
        lines += [f"#### {kind}", ""]
        lines += table(["시작 4어절", "행 수", "유형 내 비율", "검토 표시"],
                       [(phrase, count, f"{count / denominator:.2%}", "집중 반복" if count / denominator >= .1 else "") for phrase, count in entries])
    lines += ["### 반복 구문", "",
              "문장부호를 제외한 연속 2~5어절(6자 이상)의 문서 빈도입니다. 한 리뷰에서 여러 번 나와도 1회만 셉니다. 같은 빈도의 포함 구문은 가능한 한 긴 구문을 표시합니다.", ""]
    for kind, entries in result["expressions"].items():
        denominator = result["lengths"][kind]["count"]
        lines += [f"#### {kind}", ""]
        lines += table(["반복 구문", "포함 리뷰 수", "유형 내 비율"],
                       [(phrase, count, f"{count / denominator:.2%}") for phrase, count in entries])
    lines += ["### 특정 표현의 라벨 편중", "",
              "사전 지정한 표현이 포함된 리뷰 수입니다(부분 문자열 기준). 같은 행이 여러 표현에 집계됩니다. SUSPICIOUS의 50% 이상에 나타나면 다수 지배로 표시합니다. 이 목록은 모든 언어적 편향을 망라하지 않습니다.", ""]
    lines += table(["표현", "SUSPICIOUS 포함 수", "SUSPICIOUS 비율", "NORMAL 포함 수", "NORMAL 비율", "다수 지배"],
                   [(p, s, f"{s / st:.2%}" if st else "0%", n, f"{n / nt:.2%}" if nt else "0%", "예" if st and s / st >= .5 else "아니오")
                    for p, s, st, n, nt in result["probes"]])
    lines += ["### 해석과 학습 전 주의", "",
              "- exact unique 여부는 의미·문체 다양성을 보장하지 않습니다. 서로 다른 문장도 같은 생성 골격과 표현을 공유할 수 있습니다.",
              "- 조합형 synthetic은 같은 시작 구문·연결 표현·확신형 문체를 반복합니다. 위 빈도와 50개 샘플을 보고 사람이 자연스러움 및 타입 적합성을 검토해야 합니다.",
              "- COPY_VARIATION의 문구 반복은 설계상 의도된 특성이지만, 다른 유형의 표현 편중까지 정당화하지는 않습니다.",
              "- 대부분의 SUSPICIOUS가 synthetic이므로, 라벨 편중 표현은 실제 위험 신호가 아니라 생성 방식의 흔적일 수 있습니다. 이 audit은 데이터 변경이나 학습 적합성 인증이 아닙니다.",
              "- 실제 SUSPICIOUS는 1건이며 train에만 존재합니다. 실제 평가셋 positive recall을 검증할 수 없고 synthetic 성능을 실제 성능으로 해석하면 안 됩니다.", "",
              "## 4. label_origin × split", ""]
    origins = sorted({r["label_origin"] for r in rows})
    lines += table(["label_origin", *SPLITS, "합계"],
                   [(origin, *(result["label_origin_x_split"].get((origin, split), 0) for split in SPLITS),
                     sum(result["label_origin_x_split"].get((origin, split), 0) for split in SPLITS)) for origin in origins])
    lines += ["## 5. platform × label × split", ""]
    lines += table(["platform", "label", *SPLITS],
                   [(platform, label, *(result["platform_x_label_x_split"].get((platform, label, split), 0) for split in SPLITS))
                    for platform in sorted({r["platform"] for r in rows}) for label in ("NORMAL", "SUSPICIOUS")])
    lines += ["## 6. Leakage 재검증", ""]
    lines += table(["검사", "결과"], result["leakage"].items())
    lines += ["group/parent/family의 split 중복 및 실제 부모 행과 자식의 group/split/ID 일치를 확인하고 기존 validate도 재실행했습니다. 유사도 그래프 자체를 재학습하거나 새로운 그룹을 만들지는 않았습니다.", "",
              "## 7. 재현 가능한 사람 검토 샘플", "",
              "각 모집단을 dataset Excel 행 순서로 정렬한 뒤 seed=42의 단일 RNG로 유형 순서대로 비복원 추출했습니다. 원문은 HTML escape만 적용해 그대로 표시하며 라벨을 변경하지 않았습니다.", ""]
    for category, members in samples.items():
        lines += [f"### {category} — {len(members)}개", ""]
        for row in members:
            lines += [f"#### Excel 행 {row['dataset_excel_row']}", "",
                      f"- origin: {row['label_origin']} / label: {row['label']} / platform: {escaped(row['platform'])} / split: {row['split']}",
                      "", "<pre>" + html.escape(row["content"]) + "</pre>", ""]
    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser(description="데이터셋 읽기 전용 audit; Markdown만 생성")
    parser.add_argument("--input", type=Path, default=PROJECT_ROOT / "outputs/ptext_training_dataset_v1.xlsx")
    parser.add_argument("--report", type=Path, default=PROJECT_ROOT / "reports/ptext_training_dataset_v1_audit.md")
    args = parser.parse_args()
    source, report = args.input.resolve(), args.report.resolve()
    if source == report or report.exists():
        raise ValueError("입력과 보고서 경로가 같거나 보고서가 이미 있습니다. 덮어쓰지 않습니다.")
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    rows = read_dataset(source)
    result, repeated, samples = audit(rows, seed=42)
    content = render_report(rows, result, repeated, samples, before, seed=42)
    if hashlib.sha256(source.read_bytes()).hexdigest() != before:
        raise ValueError("입력 데이터셋 해시 변경 감지")
    report.parent.mkdir(parents=True, exist_ok=True)
    with report.open("x", encoding="utf-8") as stream:
        stream.write(content)
    print(json.dumps({key: result[key] for key in ("total", "conflict_groups", "conflict_rows", "conflict_pairs", "exact", "lengths", "leakage", "probes")}, ensure_ascii=False, indent=2))
    print("samples:", {key: len(value) for key, value in samples.items()})
    print("source SHA256 unchanged:", before)
    print("report:", report)


if __name__ == "__main__":
    main()
