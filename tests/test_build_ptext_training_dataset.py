"""Dataset assembly tests use artificial workbooks in tmp_path only."""

from collections import Counter
from copy import deepcopy
from pathlib import Path
import random
import sys

from openpyxl import Workbook, load_workbook
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import build_ptext_training_dataset as training
from generate_ptext_synthetic import TYPES, generate


@pytest.fixture
def inputs(tmp_path):
    rng = random.Random(7)
    raw = []
    for index in range(90):
        # Distinct artificial contexts prevent accidental all-to-all similarity.
        content = "".join(chr(rng.randrange(0xAC00, 0xD7A3)) for _ in range(65))
        content += " 청소솔 디자인 손잡이" if index % 2 else " 크림 보습 피부"
        raw.append([("ohouse", "oliveyoung", "elevenst", "musinsa")[index % 4],
                    index + .5, index + 100, content, 4, None, None, None, None, None, None])
    raw[20][3] = raw[0][3]  # Non-candidate exact duplicate of a reviewed NORMAL.
    reviews = tmp_path / "reviews.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reviews"
    sheet.append(list(training.read_reviews.__globals__["COLUMNS"]))
    for row in raw:
        sheet.append(row)
    workbook.save(reviews)
    workbook.close()
    labels = tmp_path / "labels.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "candidates"
    sheet.append(["source_excel_row", "platform", "product_id", "review_id", "content", "human_label"])
    for index, label in enumerate(["NORMAL"] * 9 + ["UNCERTAIN", "SUSPICIOUS"]):
        sheet.append([index + 2, *raw[index][:4], label])
    workbook.save(labels)
    workbook.close()
    return reviews, labels


def test_dataset_end_to_end_and_reproducible(inputs, tmp_path):
    before = {path: path.read_bytes() for path in inputs}
    rows, summary, metadata = training.build_dataset(*inputs, sample_count=24, synthetic_count=30, seed=42)
    assert summary["total"] == 9 + 1 + 24 + 30
    assert summary["label"] == {"NORMAL": 33, "SUSPICIOUS": 31}
    assert summary["label_origin"] == {
        "human_candidate_review": 9, "human_suspicious": 1,
        "sampled_non_candidate": 24, "synthetic": 30,
    }
    assert metadata["uncertain_excluded"] == 1
    assert not any(row.get("source_excel_row") in (11, 22) for row in rows)
    assert metadata["synthetic_unique_texts"] == 30
    synthetic = [row for row in rows if row["label_origin"] == "synthetic"]
    assert Counter(row["synthetic_type"] for row in synthetic) == {kind: 6 for kind in TYPES}
    assert all(isinstance(row["parent_review_id"], int) for row in synthetic)
    assert all(isinstance(row["product_id"], float) for row in rows)
    assert all(isinstance(row["review_id"], int) for row in rows if row["label_origin"] != "synthetic")
    assert all(row["group_id"] for row in rows)
    assert not any(metadata["leakage"].values())
    assert set(row["label"] for row in rows) == {"NORMAL", "SUSPICIOUS"}
    assert all(row["synthetic_type"] is None for row in rows if row["label_origin"] != "synthetic")
    again, again_summary, again_metadata = training.build_dataset(*inputs, sample_count=24, synthetic_count=30, seed=42)
    assert (rows, summary, metadata) == (again, again_summary, again_metadata)
    assert all(path.read_bytes() == content for path, content in before.items())
    output, report = tmp_path / "dataset.xlsx", tmp_path / "summary.md"
    training.save_outputs(rows, summary, metadata, {}, output, report)
    workbook = load_workbook(output, read_only=True)
    try:
        assert workbook.sheetnames == ["dataset", "summary"]
        records = list(workbook["dataset"].values)
        assert len(records) == summary["total"] + 1
        assert records[0] == training.COLUMNS
        assert isinstance(records[1][1], float)
        assert isinstance(records[1][2], int)
    finally:
        workbook.close()
    assert "실제 positive recall" in report.read_text(encoding="utf-8")
    assert "Real-data evaluation set" in report.read_text(encoding="utf-8")
    assert "Synthetic challenge set" in report.read_text(encoding="utf-8")
    old = output.read_bytes()
    with pytest.raises(ValueError, match="덮어쓰지"):
        training.save_outputs(rows, summary, metadata, {}, output, report)
    assert output.read_bytes() == old


def test_uncertain_never_seed_or_training(inputs):
    rows, _, _ = training.build_dataset(*inputs, sample_count=10, synthetic_count=10)
    assert all(row.get("source_excel_row") != 11 for row in rows)
    assert all(row.get("parent_source_excel_row") != 11 for row in rows)


def test_stratified_sampling_cap():
    pool = [{"platform": platform, "source_excel_row": i + offset}
            for platform, size, offset in (("ohouse", 100, 0), ("oliveyoung", 100, 100), ("musinsa", 30, 200))
            for i in range(size)]
    sample, quotas = training.stratified_sample(pool, 100, seed=42, musinsa_cap=5)
    assert len(sample) == 100
    assert quotas["musinsa"] == 5
    assert abs(quotas["ohouse"] - quotas["oliveyoung"]) <= 1
    assert training.stratified_sample(pool, 100, seed=42, musinsa_cap=5) == (sample, quotas)


def test_near_and_transitive_grouping():
    text = "실제로 사용해 보니 손잡이가 편하고 구석의 먼지도 쉽게 제거할 수 있었습니다."
    contents = [text, text.replace("쉽게", "아주 쉽게"), text.replace(" ", "\n"), "다른 내용"]
    groups = training.Groups(len(contents))
    training.union_similar(groups, contents, .7)
    assert groups.find(0) == groups.find(1) == groups.find(2)
    assert groups.find(3) != groups.find(0)


def test_human_normals_distributed_when_groups_allow(inputs):
    rows, _, _ = training.build_dataset(*inputs, sample_count=24, synthetic_count=0)
    splits = {row["split"] for row in rows if row["label_origin"] == "human_candidate_review"}
    assert splits == {"train", "validation", "test"}


def test_parent_leakage_rejected(inputs):
    rows, _, _ = training.build_dataset(*inputs, sample_count=10, synthetic_count=10)
    changed = deepcopy(rows)
    child = next(row for row in changed if row["label_origin"] == "synthetic")
    child["split"] = "test" if child["split"] != "test" else "train"
    with pytest.raises(ValueError, match="leakage"):
        training.validate(changed, {})


@pytest.mark.parametrize("field,value", [("content", "  "), ("label", "UNCERTAIN"), ("synthetic_type", "UNKNOWN")])
def test_validation_rejects_invalid_training_rows(inputs, field, value):
    rows, _, _ = training.build_dataset(*inputs, sample_count=5, synthetic_count=5)
    child = next(row for row in rows if row["label_origin"] == "synthetic")
    child[field] = value
    with pytest.raises(ValueError):
        training.validate(rows, {})


def test_label_file_must_be_valid_and_match_source(inputs):
    reviews, labels = inputs
    workbook = load_workbook(labels)
    workbook["candidates"]["E2"] = "원문 불일치"
    workbook.save(labels)
    workbook.close()
    with pytest.raises(ValueError, match="대조 실패"):
        training.build_dataset(reviews, labels, 0, 0)


def test_empty_label_file_fails_closed(inputs):
    reviews, labels = inputs
    labels.write_bytes(b"")
    with pytest.raises(ValueError, match="0바이트"):
        training.build_dataset(reviews, labels)


def test_generator_types_parent_families_and_no_metadata_in_text():
    parent = {"source_excel_row": 2, "platform": "platform_sentinel", "product_id": 123,
              "review_id": 456, "content": "청소솔 손잡이 디자인이 편하고 청소할 때 물기 제거가 쉬웠어요."}
    rows = generate([parent], 30, seed=42)
    assert rows == generate([parent], 30, seed=42)
    assert set(row["synthetic_type"] for row in rows) == set(TYPES)
    assert all(row["parent_review_id"] == 456 for row in rows)
    assert all(row["label_origin"] == "synthetic" for row in rows)
    assert all("platform_sentinel" not in row["content"] for row in rows)
    assert all(not any(kind in row["content"] for kind in TYPES) for row in rows)
    copies = [row for row in rows if row["synthetic_type"] == "COPY_VARIATION"]
    assert sorted(Counter(row["synthetic_family_id"] for row in copies).values()) == [3, 3]
